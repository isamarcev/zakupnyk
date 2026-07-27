"""Parser for the needs-list Excel format.

Expected layout per sheet (header in row 1):
    A: Найменування  C: Потреба  D: придбано/замовлено  E: в роботі
    F: ціна  I: специфікація
    J/K: Prom.ua (url, price)  L/M: rozetka (url, price)  N/O: дистриб'ютор (url, price)

Rows without a name or with a "Загальна вартість" summary line are skipped.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from procflow.application.dto import ImportedOffer, ImportedRow, ImportedSheet
from procflow.domain.entities import SupplierSource
from procflow.domain.exceptions import XlsxImportError

_SUMMARY_MARKERS = ("загальна вартість", "усього", "всього")
_OFFER_COLUMNS: list[tuple[int, int, SupplierSource]] = [
    (9, 10, SupplierSource.PROM),  # J, K (0-based)
    (11, 12, SupplierSource.ROZETKA),  # L, M
    (13, 14, SupplierSource.DISTRIBUTOR),  # N, O
]


def _as_float(value: Any) -> float | None:
    if isinstance(value, int | float):
        return float(value)
    if isinstance(value, str):
        cleaned = value.replace("\xa0", "").replace(" ", "").replace(",", ".")
        try:
            return float(cleaned)
        except ValueError:
            return None
    return None


def _as_text(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def parse_needs_xlsx(path: Path) -> list[ImportedSheet]:
    """Parse the workbook into import DTOs. Raises XlsxImportError on failure."""
    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:  # openpyxl raises various types here
        raise XlsxImportError(f"Не вдалося відкрити файл: {exc}") from exc

    sheets: list[ImportedSheet] = []
    for worksheet in workbook.worksheets:
        rows: list[ImportedRow] = []
        for raw in worksheet.iter_rows(min_row=2, values_only=True):
            parsed = _parse_row(raw)
            if parsed is not None:
                rows.append(parsed)
        if rows:
            sheets.append(ImportedSheet(title=worksheet.title.strip(), rows=rows))
    workbook.close()

    if not sheets:
        raise XlsxImportError("У файлі не знайдено жодного рядка з позиціями")
    return sheets


def _parse_row(raw: tuple[Any, ...]) -> ImportedRow | None:
    name = _as_text(raw[0] if len(raw) > 0 else None)
    if not name or any(marker in name.lower() for marker in _SUMMARY_MARKERS):
        return None
    qty_required = _as_float(raw[2]) if len(raw) > 2 else None
    if qty_required is None:
        return None

    offers = []
    for url_col, price_col, source in _OFFER_COLUMNS:
        url = _as_text(raw[url_col]) if len(raw) > url_col else ""
        price = _as_float(raw[price_col]) if len(raw) > price_col else None
        if url or price is not None:
            offers.append(ImportedOffer(source=source, url=url, price=price))

    return ImportedRow(
        name=name,
        specification=_as_text(raw[8]) if len(raw) > 8 else "",
        qty_required=qty_required,
        qty_ordered=_as_float(raw[3]) or 0.0 if len(raw) > 3 else 0.0,
        qty_in_progress=_as_float(raw[4]) or 0.0 if len(raw) > 4 else 0.0,
        unit_price=_as_float(raw[5]) if len(raw) > 5 else None,
        offers=offers,
    )
